import openpyxl
import os

file = "bank.xlsx"

# Create Excel file if it does not exist
if not os.path.exists(file):
    wb = openpyxl.Workbook()
    ws = wb.active

    # Excel column headings
    ws.append(["Account Number", "Name", "Mobile", "Age", "Balance"])

    wb.save(file)


def menu():
    print("\n===== BANK SYSTEM =====")
    print("1. Create Account")
    print("2. Deposit Money")
    print("3. Withdraw Money")
    print("4. View Balance")
    print("5. Exit")

    choice = int(input("Enter your choice: "))

    if choice == 1:
        create_account()

    elif choice == 2:
        deposit()

    elif choice == 3:
        withdraw()

    elif choice == 4:
        view_balance()

    elif choice == 5:
        print("Thank you!")

    else:
        print("Invalid choice")
        menu()


# Create Account
def create_account():
    print("\n----- Create Account -----")

    account = input("Enter account number: ")
    name = input("Enter name: ")
    mobile = input("Enter mobile number: ")
    age = int(input("Enter age: "))
    balance = float(input("Enter initial balance: "))

    wb = openpyxl.load_workbook(file)
    ws = wb.active

    # Store account information in Excel
    ws.append([account, name, mobile, age, balance])

    wb.save(file)

    print("Account created successfully!")
    print("Data saved in bank.xlsx")

    menu()


# Deposit Money
def deposit():
    print("\n----- Deposit Money -----")

    account = input("Enter account number: ")
    amount = float(input("Enter amount to deposit: "))

    wb = openpyxl.load_workbook(file)
    ws = wb.active

    for row in range(2, ws.max_row + 1):

        if str(ws.cell(row, 1).value) == account:

            balance = float(ws.cell(row, 5).value)

            balance = balance + amount

            ws.cell(row, 5).value = balance

            wb.save(file)

            print("Money deposited successfully!")
            print("Current balance:", balance)

            menu()
            return

    print("Account not found")
    menu()


# Withdraw Money
def withdraw():
    print("\n----- Withdraw Money -----")

    account = input("Enter account number: ")
    amount = float(input("Enter amount to withdraw: "))

    wb = openpyxl.load_workbook(file)
    ws = wb.active

    for row in range(2, ws.max_row + 1):

        if str(ws.cell(row, 1).value) == account:

            balance = float(ws.cell(row, 5).value)

            if amount > balance:
                print("Insufficient balance")
                print("Current balance:", balance)

            else:
                balance = balance - amount

                ws.cell(row, 5).value = balance

                wb.save(file)

                print("Money withdrawn successfully!")
                print("Remaining balance:", balance)

            menu()
            return

    print("Account not found")
    menu()


# View Balance
def view_balance():
    print("\n----- View Balance -----")

    account = input("Enter account number: ")

    wb = openpyxl.load_workbook(file)
    ws = wb.active

    for row in range(2, ws.max_row + 1):

        if str(ws.cell(row, 1).value) == account:

            name = ws.cell(row, 2).value
            balance = ws.cell(row, 5).value

            print("Name:", name)
            print("Current balance:", balance)

            menu()
            return

    print("Account not found")
    menu()


# Start program
menu()
